import json
from docx import Document
import re
import sys


class ConfigLoader:
    @staticmethod
    def load_config(path: str):
        with open(path, 'r') as file:
            return json.load(file)

class CourseDocumentParser:
    def __init__(self, document_path, config_path):
        self.document_path = document_path
        self.document = Document(document_path)
        self.config = ConfigLoader.load_config(config_path)

        self.paragraph_mapping_keywords = self.config["paragraph_keywords"]
        self.non_paragraph_mapping_keywords = self.config["non_paragraph_keywords"]
        self.all_keywords = set(self.paragraph_mapping_keywords.keys()) | set(self.non_paragraph_mapping_keywords)

    def extract_course_info(self):
        paragraph_extractor = ParagraphExtractor(self.document, self.paragraph_mapping_keywords, self.all_keywords)
        table_extractor = TableExtractor(self.document.tables)

        combined_data = {
            **paragraph_extractor.extract_paragraph_data(), #** to unpack dicts, take to dicts and combine them into one
            **table_extractor.extract_table_data()
        }

        return combined_data


class ParagraphExtractor:
    def __init__(self, document, paragraph_mapping_keywords, all_keywords):
        self.document = document
        self.paragraph_mapping_keywords = paragraph_mapping_keywords
        self.all_keywords = all_keywords

    def extract_paragraph_data(self):
        paragraphs = [p.text.strip() for p in self.document.paragraphs if p.text.strip()]
        extracted_data = {}

        extracted_data["name"] = self._extract_name(paragraphs[0])
        extracted_data["code"] = self._extract_code(paragraphs[0])

        current_field_keyword = None
        current_field_value = ""
        for i, paragraph_text in enumerate(paragraphs[1:], 1): #skip first paragraph
            field_name = self._get_matching_field(paragraph_text)

            # if field name in paragraph and not currently collecting
            if field_name:
                if ":" in paragraph_text:
                    extracted_data[field_name] = paragraph_text.split(":", 1)[1].strip()
                else:
                    current_field_keyword = field_name

            # if currently collecting
            elif current_field_keyword:
                current_field_value += f"\n{paragraph_text}"
                if self._is_next_paragraph_keyword(paragraphs, i):
                    extracted_data[current_field_keyword] = current_field_value.strip()
                    current_field_keyword, current_field_value = None, ""

        if current_field_keyword:
            extracted_data[current_field_keyword] = current_field_value.strip()

        return extracted_data

    def _extract_name(self, first_paragraph):
        return re.findall(r'[“"](.*?)[”"]', first_paragraph)[0]

    def _extract_code(self, first_paragraph):
        return re.findall(r'^[^"“]*', first_paragraph)[0].strip()

    def _get_matching_field(self, paragraph_text):
        for keyword, field_name in self.paragraph_mapping_keywords.items():
            if keyword.lower() in paragraph_text.lower():
                return field_name
        return None

    def _is_next_paragraph_keyword(self, paragraphs, index):
        if index + 1 < len(paragraphs):
            next_paragraph = paragraphs[index + 1]
            return any(keyword in next_paragraph for keyword in self.all_keywords)
        return False

class TableExtractor:
    def __init__(self, tables):
        self.tables = tables

    def extract_table_data(self):
        course_info_table = self.tables[0]
        contents_table = self.tables[1]
        assessment_table = self.tables[2]

        extracted_data = self._extract_info_course_info_table(course_info_table)
        extracted_data["contents"] = self._extract_info_content_table(contents_table)
        extracted_data["examination"] = self._extract_info_assessment_table(assessment_table)

        return extracted_data

    def _extract_info_course_info_table(self, table):
        extracted_data = {
            "semester": table.rows[0].cells[1].text,
            "required": table.rows[5].cells[1].text,
            "ects": table.rows[8].cells[1].text
        }

        session_details = table.rows[6].cells[1].text
        session_hours = self._calculate_session_hours(session_details)
        extracted_data.update(session_hours)

        return extracted_data

    def _extract_info_content_table(self, table):
        table_data = []
        headers = [cell.text.strip() for cell in table.rows[0].cells]
        sub_headers = [cell.text.strip() for cell in table.rows[1].cells]

        for row in table.rows[2:-1]:
            row_data = {}

            for i, cell in enumerate(row.cells):
                header = headers[i]
                sub_header = sub_headers[i]
                cell_text = cell.text.strip()

                if sub_header:
                    row_data.setdefault(header, {})[sub_header] = cell_text
                else:
                    row_data[header] = cell.text.strip()

            table_data.append(row_data)
        return table_data

    def _extract_info_assessment_table(self, table):
        table_data = []
        headers = [cell.text.strip() for cell in table.rows[0].cells]

        for row in table.rows[1:]:
            row_data = {}
            for i, cell in enumerate(row.cells):
                row_data[headers[i]] = cell.text.strip()

            table_data.append(row_data)
        return table_data

    def _calculate_session_hours(self, session_details):
        session_hours = {"lectureHrs": 0, "tutorialHrs": 0, "labHrs": 0} #initialize to 0 to add weekly and bi-weekly if both is given for one type

        for session in self._split_sessions_into_categories(session_details):
            sws = self._calculate_sws(session)
            session_lower = session.lower()

            if "lecture" in session_lower:
                session_hours["lectureHrs"] += sws
            elif "tutorial" in session_lower:
                session_hours["tutorialHrs"] += sws
            elif "lab" in session_lower:
                session_hours["labHrs"] += sws

        return session_hours

    def _split_sessions_into_categories(self, session_string: str) -> list:
        return re.split(r' \+ |, ', session_string)

    def _calculate_sws(self, session_string: str) -> int:
        if re.search(r'bi\s*-?\s*weekly', session_string.lower()):
            number = int(session_string.split()[0])
            return number
        else:
            number = int(session_string.split()[0])
            return number * 2


if __name__ == '__main__':
    if len(sys.argv) < 2:
        print("Usage: python main.py <document_path>")
        sys.exit(1)

    document_path = sys.argv[1]
    output_path = 'output/output.json'

    parser = CourseDocumentParser(document_path, "config.json")
    course_info = parser.extract_course_info()

    with open(output_path, 'w', encoding='utf-8') as json_file:
        json.dump(course_info, json_file, indent=4)

    print(json.dumps(course_info, indent=4))
    print(f"Course information extracted successfully! Saved to {output_path}")