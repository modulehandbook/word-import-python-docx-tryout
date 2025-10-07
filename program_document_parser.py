import re
from docx import Document
from copy import deepcopy
from io import BytesIO

from course_document_parser import CourseDocumentParser

class ProgramDocumentParser:
    def __init__(self, document_path, config_path):
        self.document_path = document_path
        self.document = Document(document_path)
        self.config_path = config_path

    def extract_program_info(self):
        program_info = self._add_program_metadata()
        program_info["courses"] = []

        courses = self._split_document_into_courses()

        for course in courses:
            temp_doc = self._create_temp_doc(course)
            parser = CourseDocumentParser(temp_doc, self.config_path, raw_mode=True)
            course_data = parser.extract_course_info()
            program_info["courses"].append(course_data)

        return program_info

    def _add_program_metadata(self):
        program_metadata = {
            "name": "Informatics and Computer Science",
            "code": "INCS",
            "degree": "Bachelor",
        }
        return program_metadata

    def _split_document_into_courses(self):
        """
        Splits the document into course blocks as lists of elements (paragraphs + tables)
        without creating temp Document objects yet.
        """
        course_header_pattern = re.compile(r'^[A-Z]{2,}\s*\d+\s*[“"(].+[”")]')
        courses = []
        current_elements = []
        started = False

        for block in self.document.element.body:
            text = block.text.strip() if block.tag.endswith("p") else ""

            if text and course_header_pattern.match(text):
                # finish previous course if any
                if current_elements:
                    courses.append(current_elements)
                    current_elements = []
                started = True  # now start collecting course elements (skip everything before first header)

            if started:
                current_elements.append(block)

        if current_elements:
            courses.append(current_elements)

        return courses

    def _create_temp_doc(self, elements):
        """
        Creates a proper Document object containing the given XML elements.
        Saving and reloading ensures .tables and .paragraphs are correctly built.
        """
        temp_doc = Document()

        # Remove default paragraph
        for p in temp_doc.paragraphs:
            p._element.getparent().remove(p._element)

        body = temp_doc._element.body

        for element in elements:
            # element is already an XML object; no need for ._element
            body.append(deepcopy(element))

        # Save to in-memory file and reload to rebuild .tables
        from io import BytesIO
        temp_bytes = BytesIO()
        temp_doc.save(temp_bytes)
        temp_bytes.seek(0)
        return Document(temp_bytes)


